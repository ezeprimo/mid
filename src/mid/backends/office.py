"""Office backend — explicit opt-in COM conversion via installed Word/Excel.

v1 scope (maintainer decision): Word/Excel ONLY (``.doc``/``.xls``). ``.ppt``
stays on the LegacyPlaceholder migrate-first path, deferred post-v1.

Kill scope (maintainer decision): PID-scoped via GetWindowThreadProcessId.
``taskkill /IM`` constrained to the allowlisted Office exes is fallback only
when no PID could be resolved. The killer never executes a user-controlled
path — only ``taskkill`` with ``shell=False``.
"""

from __future__ import annotations

import gc
import os
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import ClassVar

from mid.backends.base import Availability, Backend
from mid.models import ConvertResult

#: v1 ProgIDs: Word + Excel only. PowerPoint deferred post-v1.
#: SaveAs targets are macro-free OOXML, read natively by MarkItDown:
#: 12 = wdFormatXMLDocument (*.docx), 51 = xlOpenXMLWorkbook (*.xlsx).
#: Both constants exist since Office 2007, so Office 2013 needs no fallback.
#: A macro-carrying source (.doc/.xls with VBA) saves macro-stripped under
#: DisplayAlerts=0; FileFormat 52 (.xlsm) is NOT a fallback — MarkItDown's
#: XlsxConverter only accepts the `.xlsx` extension.
PROGIDS: dict[str, tuple[str, str, int]] = {
    ".doc": ("Word.Application", "WINWORD.EXE", 12),
    ".xls": ("Excel.Application", "EXCEL.EXE", 51),
}

#: Intermediate extension per input extension. The temp file MUST carry the
#: real OOXML suffix — MarkItDown dispatches on extension.
_INTERMEDIATE_EXT: dict[str, str] = {".doc": ".docx", ".xls": ".xlsx"}

#: Only these exes may ever appear in a kill fallback argv. Never a user path.
_ALLOWED_EXES = frozenset({"WINWORD.EXE", "EXCEL.EXE"})

#: 20 MB intermediate cap shared with the LibreOffice tail.
_MAX_OUTPUT_BYTES = 20 * 1024 * 1024

#: Injectable seams for tests. ``None`` selects the real implementation.
_COM_FACTORY = None
_KILLER = None


def _default_killer(target: int | str) -> None:
    """Kill an orphaned Office process. PID-scoped first, allowlist fallback.

    Args:
        target: PID (int) from GetWindowThreadProcessId, or an exe name that
            MUST be in :data:`_ALLOWED_EXES`. Anything else is ignored —
            user-controlled paths are never executed.
    """
    if isinstance(target, int):
        argv = ["taskkill", "/F", "/PID", str(target)]
    else:
        exe = str(target).upper()
        if exe not in _ALLOWED_EXES:
            return
        argv = ["taskkill", "/F", "/IM", exe]
    try:
        subprocess.run(argv, shell=False, capture_output=True, text=True, timeout=30)
    except Exception:
        pass


def _get_killer():
    return _KILLER if _KILLER is not None else _default_killer


def _get_com_factory():
    if _COM_FACTORY is not None:
        return _COM_FACTORY

    def _factory(progid: str):
        from win32com.client import DispatchEx

        return DispatchEx(progid)

    return _factory


def _read_app_paths(exe_name: str) -> Path | None:
    """Resolve an Office exe via App Paths registry keys. Never raises."""
    try:
        import winreg

        for hive in (winreg.HKEY_LOCAL_MACHINE, winreg.HKEY_CURRENT_USER):
            try:
                with winreg.OpenKey(hive, r"SOFTWARE\Microsoft\Windows\CurrentVersion\App Paths\\" + exe_name) as key:
                    value, _ = winreg.QueryValueEx(key, "")
                    p = Path(value)
                    if p.is_file():
                        return p
            except OSError:
                continue
    except Exception:
        pass
    return None


def _read_curver(progid: str) -> str | None:
    """Read HKCR\\<ProgId>\\CurVer version token. Never raises."""
    try:
        import winreg

        with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, progid + r"\CurVer") as key:
            value, _ = winreg.QueryValueEx(key, "")
            return str(value) if value else None
    except Exception:
        return None


def _resolve_tool_path() -> tuple[Path | None, str | None]:
    """Resolve the Office install signal.

    Precedence: ``MID_OFFICE_PATH`` (must be an existing file, invalid falls
    through) else App Paths ``WINWORD.EXE`` then ``EXCEL.EXE``.

    Returns (Path|None, str|None) where the second is a hint when unresolved.
    The resolved path is a version signal only — it is never executed.
    """
    env_path = os.environ.get("MID_OFFICE_PATH")
    if env_path:
        p = Path(env_path)
        if p.is_file():
            return p, None
        # invalid env falls through to App Paths (OFFICE-02)
    for exe in ("WINWORD.EXE", "EXCEL.EXE"):
        found = _read_app_paths(exe)
        if found is not None:
            return found, None
    return None, "Office not detected — install Word/Excel 2016+ or set MID_OFFICE_PATH"


def _resolve_timeout() -> int:
    """Resolve convert timeout from MID_OFFICE_TIMEOUT, clamp 5..300 default 30."""
    raw = os.environ.get("MID_OFFICE_TIMEOUT")
    if raw is None:
        return 30
    try:
        val = int(raw)
    except (ValueError, TypeError):
        return 30
    if val < 5:
        return 5
    if val > 300:
        return 300
    return val


def _resolve_pid(app, exe: str) -> int | str:
    """Resolve the COM instance PID via GetWindowThreadProcessId.

    Returns the PID (int) when resolvable, else the allowlisted exe name so
    the kill fallback stays constrained. Never raises, never returns a user
    path.
    """
    try:
        import win32process

        hwnd = app.Hwnd
        if isinstance(hwnd, int):
            _, pid = win32process.GetWindowThreadProcessId(hwnd)
            if isinstance(pid, int) and pid > 0:
                return pid
    except Exception:
        pass
    return exe.upper() if exe.upper() in _ALLOWED_EXES else "WINWORD.EXE"


def _map_error(exc: BaseException) -> str:
    """Map COM failures to the OFFICE-05 taxonomy. Never raises."""
    try:
        msg = str(exc) or type(exc).__name__
    except Exception:
        return "conversion failed"
    # Locale-independent lock detection: WinError messages are localized
    # (e.g. Spanish WinError 32), so check the numeric code first (#BUG-2).
    winerror = getattr(exc, "winerror", None)
    if winerror in (32, 33):  # ERROR_SHARING_VIOLATION / ERROR_LOCK_VIOLATION
        return f"file locked — close Office and retry: {msg}"
    if winerror == 5:  # ERROR_ACCESS_DENIED
        return f"permission denied — close Office and retry: {msg}"
    lowered = msg.lower()
    if "rpc_e_servercall_retrylater" in lowered or "call was rejected" in lowered or "busy" in lowered:
        return f"Office busy — retry: {msg}"
    if "locked" in lowered or "sharing violation" in lowered or "permission denied" in lowered:
        return f"file locked — close Office and retry: {msg}"
    return msg


#: Whole-cell ``NaN`` (pandas empty merged cells) in pipe context only, so
#: legitimate words containing "nan" (e.g. ``financiero``) are untouched.
_NAN_CELL_RE = re.compile(r"(?<=\|)\s*nan\s*(?=\|)", re.IGNORECASE)
#: Pandas default header for blank columns. Case-sensitive by design: only
#: the exact ``Unnamed: N`` artifact is blanked, never real header text.
_UNNAMED_CELL_RE = re.compile(r"(?<=\|)\s*Unnamed:\s*\d+\s*(?=\|)")
#: Collapse horizontal whitespace runs but preserve newlines.
_MULTI_SPACE_RE = re.compile(r"[^\S\n]{2,}")


def _normalize_ooxml_markdown(text: str) -> str:
    """Normalize OOXML-path markdown artifacts. Never raises.

    1. Literal U+00A0 → regular space, then collapse horizontal runs
       (newlines preserved).
    2. Whole-cell ``NaN`` (case-insensitive, pipe-delimited) → empty cell.
       Substrings such as ``financiero`` are untouched.
    3. Exact ``Unnamed: N`` header cells → empty cell, keeping the pipe
       count (column alignment) stable.

    Applies to the final markdown of both ``.doc`` and ``.xls`` paths;
    ``.doc`` output is already clean so it passes through unchanged.
    """
    try:
        text = text.replace("\u00a0", " ")
        text = _NAN_CELL_RE.sub("", text)
        text = _UNNAMED_CELL_RE.sub("", text)
        text = _MULTI_SPACE_RE.sub(" ", text)
        return text
    except Exception:
        return text


def _expand_merged_cells(path: Path) -> None:
    """Forward-fill Excel merged ranges so every row is self-contained. Never raises.

    For each worksheet and each merged range, the top-left value is copied
    into every cell of the range (horizontal and vertical), then the range is
    unmerged so downstream readers (pandas ``read_excel`` via MarkItDown, which
    reports non-top-left merged cells as ``NaN``) see the repeated value.
    Only cells inside a real merged range are touched — genuinely empty cells
    and all-empty spacer rows stay empty. A merged range whose top-left is
    empty is a no-op. Runs BEFORE MarkItDown; the existing
    :func:`_normalize_ooxml_markdown` still runs AFTER (no ``NaN`` /
    ``Unnamed: N`` / NBSP may be reintroduced here — only stored values move).

    Method note: the precise ``merged_cells`` mask is used instead of a blind
    pandas ``ffill`` on both axes, which would bleed prior values into spacer
    rows and unrelated blanks. If ``openpyxl`` is unavailable (or anything
    fails), the file is left untouched and conversion falls back to the
    pre-existing behavior.
    """
    try:
        import openpyxl  # lazy so a missing extra never breaks convert

        wb = openpyxl.load_workbook(path)
        try:
            for ws in wb.worksheets:
                for rng in list(ws.merged_cells.ranges):
                    top_value = ws.cell(row=rng.min_row, column=rng.min_col).value
                    ws.unmerge_cells(str(rng))
                    if top_value is None:
                        continue
                    for row in ws.iter_rows(
                        min_row=rng.min_row,
                        max_row=rng.max_row,
                        min_col=rng.min_col,
                        max_col=rng.max_col,
                    ):
                        for cell in row:
                            cell.value = top_value
            wb.save(path)
        finally:
            try:
                wb.close()
            except Exception:
                pass
    except Exception:
        pass


def _convert_inner(progid: str, src_copy: Path, out_path: Path, holder: dict, save_format: int) -> None:
    """Run the COM conversion on the worker thread. Records app for Quit."""
    try:
        try:
            import pythoncom

            pythoncom.CoInitialize()
        except Exception:
            pass
        factory = _get_com_factory()
        app = factory(progid)
        holder["app"] = app
        # Harden: hidden, no alerts/macros/DDE/links, never silent user attach.
        try:
            app.Visible = False
        except Exception:
            pass
        for attr, value in (("DisplayAlerts", 0), ("AutomationSecurity", 3), ("AskToUpdateLinks", False)):
            try:
                setattr(app, attr, value)
            except Exception:
                pass
        is_excel = progid == "Excel.Application"
        try:
            if is_excel:
                wb = app.Workbooks.Open(
                    str(src_copy),
                    ReadOnly=True,
                    AddToMru=False,
                )
                try:
                    wb.SaveAs(str(out_path), FileFormat=save_format)
                finally:
                    try:
                        wb.Close(False)
                    except Exception:
                        pass
            else:
                # NOTE (#28): Documents.Open has no WithWindow parameter — Word 2013
                # rejects the unknown kwarg. Hidden mode is already enforced via
                # app.Visible = False above, which is supported on all versions.
                docs = app.Documents.Open(
                    str(src_copy),
                    ReadOnly=True,
                    AddToRecentFiles=False,
                )
                try:
                    docs.SaveAs(str(out_path), FileFormat=save_format)
                finally:
                    try:
                        docs.Close(0)
                    except Exception:
                        pass
        finally:
            try:
                app.Quit()
            except Exception:
                pass
        holder["ok"] = True
    except BaseException as exc:  # never-raise; surfaced via holder
        holder["error"] = exc
    finally:
        try:
            import pythoncom

            pythoncom.CoUninitialize()
        except Exception:
            pass


class OfficeBackend(Backend):
    """Explicit opt-in COM backend for legacy Word/Excel on Windows."""

    name: str = "office"
    display_name: str = "Office"
    opt_in: bool = True
    supported_extensions: ClassVar[frozenset[str]] = frozenset({".doc", ".xls"})
    required_tools: ClassVar[tuple[str, ...]] = ("WINWORD.EXE", "EXCEL.EXE")

    def probe(self) -> Availability:
        """Probe availability — never raises, completes within 2s."""
        start = time.monotonic()
        try:
            if sys.platform != "win32":
                return Availability(
                    available=False,
                    version=None,
                    tool_path=None,
                    reason="unsupported platform: Office backend requires Windows (win32)",
                    checked_at=datetime.now(timezone.utc),
                )
            tool_path, hint = _resolve_tool_path()
            if tool_path is None:
                return Availability(
                    available=False,
                    version=None,
                    tool_path=None,
                    reason=hint or "Office not detected",
                    checked_at=datetime.now(timezone.utc),
                )
            if time.monotonic() - start > 2:
                return Availability(
                    available=False,
                    version=str(tool_path),
                    tool_path=str(tool_path),
                    reason="probe timed out",
                    checked_at=datetime.now(timezone.utc),
                )
            try:
                from win32com.client import DispatchEx  # noqa: F401
            except Exception:
                return Availability(
                    available=False,
                    version=None,
                    tool_path=str(tool_path),
                    reason="pywin32 not installed — install the office-windows extra",
                    checked_at=datetime.now(timezone.utc),
                )
            version = _read_curver("Word.Application") or _read_curver("Excel.Application")
            return Availability(
                available=True,
                version=version,
                tool_path=str(tool_path),
                reason=None,
                checked_at=datetime.now(timezone.utc),
            )
        except Exception as exc:  # never-raise
            return Availability(
                available=False,
                version=None,
                tool_path=None,
                reason=str(exc) or "probe failed",
                checked_at=datetime.now(timezone.utc),
            )

    def convert(self, path: Path) -> ConvertResult:
        """Convert via per-app DispatchEx -> OOXML -> MarkItDown. Never raises.

        ``.doc`` saves as ``.docx`` (12 = wdFormatXMLDocument), ``.xls`` as
        ``.xlsx`` (51 = xlOpenXMLWorkbook); the intermediate is then read by
        MarkItDown's native docx/xlsx reader — no HTML involved. MarkItDown's
        XlsxConverter iterates every sheet (``## <name>`` heading per sheet),
        so multi-sheet workbooks convert whole. On the ``.xls`` path merged
        ranges are forward-filled first (top-left value repeated into every
        merged cell, ranges unmerged) so each markdown row is self-contained
        for AI readers; spacer rows stay empty. Final markdown is normalized
        (NBSP → space, whole-cell NaN → empty, ``Unnamed: N`` → empty).
        """
        try:
            if not path.is_file():
                return ConvertResult(content="", metadata={}, success=False, error=f"file not found: {path}")
            ext = path.suffix.lower()
            if ext not in self.supported_extensions:
                return ConvertResult(content="", metadata={}, success=False, error=f"unsupported format {ext}")
            if sys.platform != "win32":
                return ConvertResult(
                    content="",
                    metadata={},
                    success=False,
                    error="unsupported platform: Office backend requires Windows (win32)",
                )
            progid, exe, save_format = PROGIDS[ext]
            timeout = _resolve_timeout()

            with tempfile.TemporaryDirectory(prefix="mid-office-") as tmpdir:
                tmp = Path(tmpdir)
                tmp_input = tmp / path.name
                try:
                    # NOTE: the copy runs BEFORE COM opens anything, so a file
                    # merely open in Office (shared-read) still converts — the
                    # live instance is never attached to. Only an exclusive
                    # lock fails here, mapped locale-independently (#BUG-2).
                    shutil.copy2(path, tmp_input)
                except OSError as exc:
                    return ConvertResult(content="", metadata={}, success=False, error=_map_error(exc))
                intermediate = tmp / f"{path.stem}{_INTERMEDIATE_EXT[ext]}"

                holder: dict = {}
                worker = threading.Thread(
                    target=_convert_inner,
                    args=(progid, tmp_input, intermediate, holder, save_format),
                    daemon=True,
                )
                worker.start()
                worker.join(timeout)
                # Deterministic COM release (#29): pop the app out of the holder
                # immediately. The stored wrapper plus traceback->frame cycles
                # otherwise keep EXCEL.EXE/WINWORD.EXE alive nondeterministically
                # (CoUninitialize already ran while wrappers were still alive).
                # The finally below drains reference cycles synchronously so the
                # Office process can exit before we return.
                app = holder.pop("app", None)
                try:
                    if worker.is_alive():
                        # Hang: Quit, then PID-scoped kill fallback, tmp auto-removed.
                        if app is not None:
                            try:
                                app.Quit()
                            except Exception:
                                pass
                        try:
                            _get_killer()(_resolve_pid(app, exe) if app is not None else exe)
                        except Exception:
                            pass
                        return ConvertResult(
                            content="",
                            metadata={},
                            success=False,
                            error=f"conversion timed out after {timeout}s (orphan cleaned up)",
                        )
                    if "error" in holder:
                        return ConvertResult(content="", metadata={}, success=False, error=_map_error(holder["error"]))

                    try:
                        size = intermediate.stat().st_size
                    except OSError as exc:
                        return ConvertResult(content="", metadata={}, success=False, error=str(exc))
                    if size > _MAX_OUTPUT_BYTES:
                        return ConvertResult(content="", metadata={}, success=False, error="output exceeds 20 MB limit")
                    if size == 0:
                        return ConvertResult(content="", metadata={}, success=False, error="conversion produced empty output")
                    try:
                        from mid.converters.markitdown import MarkitDownConverter
                    except Exception as exc:
                        return ConvertResult(content="", metadata={}, success=False, error=str(exc))
                    if ext == ".xls":
                        # AI-legibility: expand merged headers/cells into every
                        # row before MarkItDown reads the sheet (never raises;
                        # falls back to the unfilled intermediate on failure).
                        _expand_merged_cells(intermediate)
                    try:
                        md = MarkitDownConverter()
                        md_result = md.convert(intermediate)
                        if not md_result.success:
                            return ConvertResult(
                                content="", metadata={}, success=False, error=md_result.error or "delegation failed"
                            )
                        return ConvertResult(
                            content=_normalize_ooxml_markdown(md_result.content),
                            metadata={"source": path.name, "format": ext.lstrip("."), "success": True},
                            success=True,
                            error=None,
                        )
                    except Exception as exc:
                        return ConvertResult(content="", metadata={}, success=False, error=str(exc))
                finally:
                    holder.clear()
                    del app
                    gc.collect()
        except Exception as exc:  # never-raise outer
            return ConvertResult(content="", metadata={}, success=False, error=str(exc) or "conversion failed")

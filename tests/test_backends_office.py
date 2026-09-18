"""Unit tests for OfficeBackend — strict TDD for OFFICE-01..06.

v1 scope (maintainer decision): Word/Excel ONLY (.doc/.xls). .ppt stays on
the LegacyPlaceholder migrate-first path, deferred post-v1.

Kill scope (maintainer decision): PID-scoped via GetWindowThreadProcessId.
taskkill /IM constrained to the allowlisted Office exes is fallback only.
"""

from __future__ import annotations

import os
import sys
import tempfile
import threading
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest

pytestmark = pytest.mark.filterwarnings("ignore::DeprecationWarning")


@pytest.fixture
def win32():
    """Pretend to run on Windows without touching the real platform."""
    with patch.object(sys, "platform", "win32"):
        yield


@pytest.fixture
def com_stub(monkeypatch):
    """Stub win32com.client so lazy imports succeed on Linux."""
    fake = MagicMock(name="win32com.client")
    monkeypatch.setitem(sys.modules, "win32com", MagicMock(name="win32com"))
    monkeypatch.setitem(sys.modules, "win32com.client", fake)
    return fake


def _backend():
    from mid.backends.office import OfficeBackend

    return OfficeBackend()


# --- 2.1 RED: killer argv ----------------------------------------------------

# Task 2.1 RED — killer uses list argv, shell=False, exe from allowlist only.


def test_killer_argv_pid_scoped():
    """Default killer is PID-scoped: list argv, shell=False, no user path."""
    from mid.backends.office import _default_killer

    with patch("mid.backends.office.subprocess.run") as mock_run:
        _default_killer(1234)
        assert mock_run.call_count == 1
        argv = mock_run.call_args[0][0]
        assert isinstance(argv, list)
        assert argv == ["taskkill", "/F", "/PID", "1234"]
        kwargs = mock_run.call_args.kwargs
        assert kwargs.get("shell") is False


def test_killer_argv_exe_allowlist_only(tmp_path):
    """Exe fallback is constrained to the allowlist; user paths never run."""
    from mid.backends.office import _ALLOWED_EXES, _default_killer

    assert _ALLOWED_EXES == frozenset({"WINWORD.EXE", "EXCEL.EXE"})
    evil = str(tmp_path / "evil.exe")
    with patch.dict(os.environ, {"MID_OFFICE_PATH": evil}):
        with patch("mid.backends.office.subprocess.run") as mock_run:
            _default_killer("WINWORD.EXE")
            argv = mock_run.call_args[0][0]
            assert isinstance(argv, list)
            assert kwargs_shell_false(mock_run)
            assert "WINWORD.EXE" in argv
            assert evil not in argv
            mock_run.reset_mock()
            # user-controlled path is never executed, even if requested
            _default_killer(evil)
            assert mock_run.call_count == 0
            _default_killer("outlook.exe")
            assert mock_run.call_count == 0


def kwargs_shell_false(mock_run):
    return mock_run.call_args.kwargs.get("shell") is False


# --- 2.2 RED: hang cleanup ---------------------------------------------------

# Task 2.2 RED — hung COM asserts Quit-called + killer-called + tmp removed.


def test_hang_cleanup_quit_kill_tmp(tmp_path):
    """Hung conversion: Quit called, killer called, mid-office- tmp removed."""
    import mid.backends.office as office_mod

    src = tmp_path / "in.doc"
    src.write_text("x", encoding="utf-8")

    block = threading.Event()
    app = MagicMock(name="WordApp")

    def fake_factory(progid):
        assert progid == "Word.Application"
        return app

    def hanging_saveas(html_path, *args, **kwargs):
        block.wait(timeout=30)

    docs = MagicMock()
    app.Documents.Open.return_value = docs
    docs.SaveAs.side_effect = hanging_saveas

    killer_calls = []
    before = {p.name for p in Path(tempfile.gettempdir()).iterdir() if p.name.startswith("mid-office-")}

    with patch.object(sys, "platform", "win32"):
        with patch.dict(os.environ, {"MID_OFFICE_TIMEOUT": "5"}):
            with patch.object(office_mod, "_COM_FACTORY", fake_factory):
                with patch.object(office_mod, "_KILLER", killer_calls.append):
                    b = office_mod.OfficeBackend()
                    result = b.convert(src)

    assert result.success is False
    assert "timed out" in (result.error or "").lower()
    assert app.Quit.called, "Quit must be called on hang"
    assert len(killer_calls) == 1, "kill fallback must run on hang"
    after = {p.name for p in Path(tempfile.gettempdir()).iterdir() if p.name.startswith("mid-office-")}
    assert after - before == set(), "mid-office- tmp must be removed even on timeout"


# --- 2.3 RED: bad-path fallthrough -------------------------------------------

# Task 2.3 RED — invalid MID_OFFICE_PATH falls through to App Paths.


def test_bad_path_fallthrough(win32, com_stub, tmp_path):
    """Invalid MID_OFFICE_PATH falls through to App Paths resolution."""
    from mid.backends import office as office_mod

    app_paths_exe = tmp_path / "WINWORD.EXE"
    app_paths_exe.write_text("x")
    with patch.dict(os.environ, {"MID_OFFICE_PATH": "/nonexistent/office.exe"}):
        with patch.object(office_mod, "_read_app_paths", return_value=app_paths_exe) as mock_app:
            with patch.object(office_mod, "_read_curver", return_value="16.0"):
                avail = office_mod.OfficeBackend().probe()
                assert mock_app.called, "bad env path must fall through to App Paths"
                assert avail.available is True
                assert avail.tool_path == str(app_paths_exe)


# --- OFFICE-01: identity ------------------------------------------------------


def test_office_identity_word_excel_only():
    """v1 = Word/Excel ONLY; .ppt deferred post-v1 (maintainer decision)."""
    from mid.backends.office import PROGIDS, OfficeBackend

    assert OfficeBackend.name == "office"
    assert OfficeBackend.opt_in is True
    assert OfficeBackend.supported_extensions == frozenset({".doc", ".xls"})
    assert set(PROGIDS) == {".doc", ".xls"}
    assert PROGIDS[".doc"][0] == "Word.Application"
    assert PROGIDS[".xls"][0] == "Excel.Application"
    assert ".ppt" not in OfficeBackend.supported_extensions


def test_office_registration_win32_only(win32, com_stub):
    """Guarded registration: win32-only, ValueError-safe, absent on Linux."""
    from mid.backends.registry import BackendRegistry

    reg = BackendRegistry()
    with patch.object(sys, "platform", "linux"):
        import importlib

        import mid.backends.office as office_mod

        importlib.reload(office_mod)
        assert reg.get("office") is None
    # win32 path registers without raising on duplicates
    from mid.backends.office import OfficeBackend

    reg.register(OfficeBackend())
    try:
        reg.register(OfficeBackend())
        raise AssertionError("duplicate registration must raise")
    except ValueError:
        pass


def test_office_absent_on_linux():
    """On Linux the win32 guard keeps office out of the shared registry."""
    from mid.backends.registry import registry

    names = [b.name for b in registry.list_all()]
    if sys.platform != "win32":
        assert "office" not in names
    else:
        pytest.skip("win32 runner registers office by design")


# --- OFFICE-02: probe ----------------------------------------------------------


def test_probe_platform_gate():
    """Non-win32 probe is unavailable with unsupported-platform reason."""
    with patch.object(sys, "platform", "linux"):
        avail = _backend().probe()
        assert avail.available is False
        assert "unsupported platform" in (avail.reason or "").lower()


def test_probe_env_precedence_valid(win32, com_stub, tmp_path):
    from mid.backends import office as office_mod

    fake = tmp_path / "WINWORD.EXE"
    fake.write_text("x")
    with patch.dict(os.environ, {"MID_OFFICE_PATH": str(fake)}):
        with patch.object(office_mod, "_read_curver", return_value="16.0"):
            avail = office_mod.OfficeBackend().probe()
            assert avail.available is True
            assert avail.tool_path == str(fake)
            assert avail.version == "16.0"


def test_probe_needs_win32com(win32, tmp_path):
    """Missing pywin32 maps to unavailable, never raises."""
    from mid.backends import office as office_mod

    fake = tmp_path / "WINWORD.EXE"
    fake.write_text("x")
    with patch.dict(os.environ, {"MID_OFFICE_PATH": str(fake)}):
        with patch.dict(sys.modules, {"win32com": None, "win32com.client": None}):
            sys.modules.pop("win32com.client", None)
            sys.modules.pop("win32com", None)
            with patch.object(office_mod, "_read_curver", return_value=None):
                # force ImportError by blocking the import machinery
                import builtins

                real_import = builtins.__import__

                def fake_import(name, *args, **kwargs):
                    if name.startswith("win32com"):
                        raise ImportError("No module named win32com")
                    return real_import(name, *args, **kwargs)

                with patch.object(builtins, "__import__", fake_import):
                    avail = office_mod.OfficeBackend().probe()
                    assert avail.available is False
                    assert "pywin32" in (avail.reason or "").lower()


def test_probe_never_raises(win32):
    from mid.backends import office as office_mod

    with patch.object(office_mod, "_resolve_tool_path", side_effect=RuntimeError("boom")):
        avail = office_mod.OfficeBackend().probe()
        assert avail.available is False
        assert "boom" in (avail.reason or "")


# --- OFFICE-02/04: timeout clamp ----------------------------------------------


def test_resolve_timeout_clamp():
    from mid.backends.office import _resolve_timeout

    with patch.dict(os.environ, {}, clear=True):
        assert _resolve_timeout() == 30
    with patch.dict(os.environ, {"MID_OFFICE_TIMEOUT": "abc"}):
        assert _resolve_timeout() == 30
    with patch.dict(os.environ, {"MID_OFFICE_TIMEOUT": "1"}):
        assert _resolve_timeout() == 5
    with patch.dict(os.environ, {"MID_OFFICE_TIMEOUT": "9999"}):
        assert _resolve_timeout() == 300
    with patch.dict(os.environ, {"MID_OFFICE_TIMEOUT": "60"}):
        assert _resolve_timeout() == 60


# --- OFFICE-03: explicit selection + ambiguity ---------------------------------


def test_explicit_office_selection_and_ambiguity(win32, com_stub, tmp_path):
    """backend='office' selects Office; backend=None with 2 claimants -> None."""
    from mid.backends import office as office_mod
    from mid.engine import resolve_backend

    office = office_mod.OfficeBackend()
    other = MagicMock()
    other.name = "libreoffice"
    other.supported_extensions = frozenset({".doc"})

    from mid.backends.registry import registry

    original = list(registry.list_all())
    registry.clear()
    try:
        registry.register(office)
        registry.register(other)
        assert resolve_backend(".doc", preferred="office") is office
        # ambiguity: no preference, two claimants -> None (select-one)
        assert resolve_backend(".doc") is None
    finally:
        registry.clear()
        for b in original:
            try:
                registry.register(b)
            except ValueError:
                pass


def test_unknown_backend_exits_2(tmp_path, capsys):
    """CLI maps unknown backend to exit 2."""
    import sys as _sys

    from mid.backends.registry import registry

    original = list(registry.list_all())
    registry.clear()
    try:
        f = tmp_path / "sample.doc"
        f.write_text("x", encoding="utf-8")
        from mid.cli import main

        with patch.object(_sys, "argv", ["mid", "convert", str(f), "--backend", "office"]):
            with pytest.raises(SystemExit) as exc:
                main()
            assert exc.value.code == 2
        err = capsys.readouterr().err
        assert "unknown backend" in err.lower()
    finally:
        registry.clear()
        for b in original:
            try:
                registry.register(b)
            except ValueError:
                pass


def test_unavailable_stub_backend_exits_3(tmp_path, capsys):
    """CLI maps an unavailable backend to exit 3 — deterministic, no Office needed (#30)."""
    import sys as _sys
    from datetime import datetime, timezone

    from mid.backends.base import Availability, Backend
    from mid.backends.registry import registry
    from mid.models import ConvertResult

    class UnavailableStub(Backend):
        def __init__(self):
            super().__init__()
            self.name = "stub-office"
            self.display_name = "stub-office"
            self.supported_extensions = frozenset({".doc"})
            self.required_tools = tuple()

        def probe(self) -> Availability:
            return Availability(
                available=False,
                version=None,
                tool_path=None,
                reason="Office not detected",
                checked_at=datetime.now(timezone.utc),
            )

        def convert(self, path: Path) -> ConvertResult:
            raise AssertionError("must not convert when unavailable")

    original = list(registry.list_all())
    registry.clear()
    try:
        f = tmp_path / "sample.doc"
        f.write_text("x", encoding="utf-8")
        registry.register(UnavailableStub())
        from mid.cli import main

        with patch.object(_sys, "argv", ["mid", "convert", str(f), "--backend", "stub-office"]):
            with pytest.raises(SystemExit) as exc:
                main()
            assert exc.value.code == 3
        err = capsys.readouterr().err
        assert "unavailable" in err.lower()
    finally:
        registry.clear()
        for b in original:
            try:
                registry.register(b)
            except ValueError:
                pass


def test_office_absent_reports_unavailable_exit_3(tmp_path, capsys):
    """Real backend: exit 3 only when Office is actually absent (#30)."""
    import sys as _sys

    from mid.backends import office as office_mod
    from mid.backends.registry import registry

    if office_mod.OfficeBackend().probe().available:
        pytest.skip("Office present on this runner — absent-path not applicable")
    original = list(registry.list_all())
    registry.clear()
    try:
        f = tmp_path / "sample.doc"
        f.write_text("x", encoding="utf-8")
        registry.register(office_mod.OfficeBackend())
        from mid.cli import main

        with patch.object(_sys, "argv", ["mid", "convert", str(f), "--backend", "office"]):
            with pytest.raises(SystemExit) as exc:
                main()
            assert exc.value.code == 3
        err = capsys.readouterr().err
        assert "unavailable" in err.lower()
    finally:
        registry.clear()
        for b in original:
            try:
                registry.register(b)
            except ValueError:
                pass


def test_office_present_passes_availability_gate(tmp_path, capsys):
    """Real backend: when Office is present, explicit selection passes the gate (#30)."""
    import sys as _sys

    from mid.backends import office as office_mod
    from mid.backends.registry import registry

    if not office_mod.OfficeBackend().probe().available:
        pytest.skip("Office absent on this runner — gate test not applicable")
    original = list(registry.list_all())
    registry.clear()
    try:
        f = tmp_path / "sample.doc"
        f.write_text("x", encoding="utf-8")
        registry.register(office_mod.OfficeBackend())
        from mid.cli import main

        # A dummy .doc may or may not convert, but it must NOT exit 3:
        # availability was established by the probe above.
        try:
            with patch.object(_sys, "argv", ["mid", "convert", str(f), "--backend", "office"]):
                main()
        except SystemExit as exc:
            assert exc.code != 3
    finally:
        registry.clear()
        for b in original:
            try:
                registry.register(b)
            except ValueError:
                pass


# --- OFFICE-04: convert boundary -----------------------------------------------


def _fake_word_app(allow):
    app = MagicMock(name="FakeWord")
    docs = MagicMock()
    app.Documents.Open.return_value = docs

    def _saveas(out_path, *args, **kwargs):
        from docx import Document

        doc = Document()
        doc.add_paragraph("hello")
        doc.save(str(out_path))

    docs.SaveAs.side_effect = _saveas
    allow.append((app, docs))
    return app


def _fake_excel_app(allow):
    app = MagicMock(name="FakeExcel")
    wb = MagicMock()
    app.Workbooks.Open.return_value = wb

    def _saveas(out_path, *args, **kwargs):
        from openpyxl import Workbook

        book = Workbook()
        book.active.title = "Data"
        book.active["A1"] = "hello"
        book.save(str(out_path))

    wb.SaveAs.side_effect = _saveas
    allow.append((app, wb))
    return app


def _fake_com_success(progid, allow):
    """Per-app shaped double: Word progid -> Documents, Excel progid -> Workbooks."""
    if "Excel" in progid:
        return _fake_excel_app(allow)
    return _fake_word_app(allow)


def test_convert_hardened_open_and_cleanup(tmp_path):
    """Convert opens hidden/read-only with guards off; Close+Quit in finally."""
    import mid.backends.office as office_mod

    src = tmp_path / "in.doc"
    src.write_text("x", encoding="utf-8")
    seen = []
    with patch.object(sys, "platform", "win32"):
        with patch.object(office_mod, "_COM_FACTORY", lambda progid: _fake_com_success(progid, seen)):
            with patch.object(office_mod, "_KILLER", MagicMock()) as mock_killer:
                mock_md = MagicMock()
                mock_md.convert.return_value = MagicMock(success=True, content="# hello", error=None)
                with patch("mid.converters.markitdown.MarkitDownConverter", return_value=mock_md):
                    result = office_mod.OfficeBackend().convert(src)
    assert result.success is True
    assert "# hello" in result.content
    app, docs = seen[0]
    assert app.Visible is False
    assert app.DisplayAlerts == 0
    app.Documents.Open.assert_called_once()
    open_kwargs = app.Documents.Open.call_args.kwargs
    assert open_kwargs.get("ReadOnly") is True
    assert "WithWindow" not in open_kwargs, "Documents.Open has no WithWindow param (Word 2013 rejects it, #28)"
    assert docs.SaveAs.call_count == 1
    assert docs.SaveAs.call_args.kwargs.get("FileFormat") == 12
    assert str(docs.SaveAs.call_args[0][0]).endswith(".docx")
    assert app.Workbooks.Open.call_count == 0
    assert docs.Close.called
    assert app.Quit.called
    assert not mock_killer.called, "killer must not run on success"


def test_convert_excel_uses_workbooks_and_xlsx_format(tmp_path):
    """Excel .xls uses Workbooks.Open/Close and SaveAs FileFormat 51 (.xlsx)."""
    import mid.backends.office as office_mod

    src = tmp_path / "in.xls"
    src.write_text("x", encoding="utf-8")
    seen = []
    progids = []

    def factory(progid):
        progids.append(progid)
        return _fake_com_success(progid, seen)

    with patch.object(sys, "platform", "win32"):
        with patch.object(office_mod, "_COM_FACTORY", factory):
            with patch.object(office_mod, "_KILLER", MagicMock()):
                mock_md = MagicMock()
                mock_md.convert.return_value = MagicMock(success=True, content="# hello", error=None)
                with patch("mid.converters.markitdown.MarkitDownConverter", return_value=mock_md):
                    result = office_mod.OfficeBackend().convert(src)
    assert result.success is True
    assert progids == ["Excel.Application"]
    app, wb = seen[0]
    app.Workbooks.Open.assert_called_once()
    open_kwargs = app.Workbooks.Open.call_args.kwargs
    assert open_kwargs.get("ReadOnly") is True
    assert app.Documents.Open.call_count == 0
    assert wb.SaveAs.call_count == 1
    assert wb.SaveAs.call_args.kwargs.get("FileFormat") == 51
    assert str(wb.SaveAs.call_args[0][0]).endswith(".xlsx")
    assert wb.Close.call_count == 1
    close_args, close_kwargs = wb.Close.call_args
    assert (close_args and close_args[0] in (0, False)) or close_kwargs.get("SaveChanges") in (0, False)
    assert app.Quit.called


def test_convert_uses_mid_office_prefix_and_copies_input(tmp_path):
    """Tmp dir uses mid-office- prefix; input is copied in, never opened in place."""
    import mid.backends.office as office_mod

    src = tmp_path / "in.xls"
    src.write_text("x", encoding="utf-8")
    prefixes = []
    real_tmp = tempfile.TemporaryDirectory
    from unittest.mock import MagicMock as _MM

    def spy_tmp(*args, **kwargs):
        prefixes.append(kwargs.get("prefix"))
        return real_tmp(*args, **kwargs)

    seen = []
    with patch.object(sys, "platform", "win32"):
        with patch("tempfile.TemporaryDirectory", side_effect=spy_tmp):
            with patch.object(office_mod, "_COM_FACTORY", lambda progid: _fake_com_success(progid, seen)):
                mock_md = _MM()
                mock_md.convert.return_value = _MM(success=True, content="# ok", error=None)
                with patch("mid.converters.markitdown.MarkitDownConverter", return_value=mock_md):
                    with patch("shutil.copy2", wraps=__import__("shutil").copy2) as mock_copy:
                        result = office_mod.OfficeBackend().convert(src)
    assert result.success is True
    assert prefixes and all(p == "mid-office-" for p in prefixes)
    assert mock_copy.called


def test_convert_unsupported_and_missing():
    b = _backend()
    r = b.convert(Path("/nonexistent/x.doc"))
    assert r.success is False
    assert "file not found" in (r.error or "").lower()


def test_convert_ppt_deferred_to_legacy(tmp_path):
    """.ppt is NOT claimed by office v1 (deferred post-v1)."""
    b = _backend()
    p = tmp_path / "deck.ppt"
    p.write_text("x", encoding="utf-8")
    r = b.convert(p)
    assert r.success is False
    assert "unsupported format" in (r.error or "").lower()


def test_intermediate_format_constants():
    """OOXML intermediate: .doc -> docx/12, .xls -> xlsx/51 (macro-free)."""
    from mid.backends.office import _INTERMEDIATE_EXT, PROGIDS

    assert PROGIDS[".doc"] == ("Word.Application", "WINWORD.EXE", 12)
    assert PROGIDS[".xls"] == ("Excel.Application", "EXCEL.EXE", 51)
    assert _INTERMEDIATE_EXT == {".doc": ".docx", ".xls": ".xlsx"}


def test_html_helpers_removed():
    """HTML-intermediate helpers are dead code after the OOXML switch."""
    import mid.backends.office as office_mod

    for name in (
        "_decode_html_bytes",
        "_detect_charset",
        "_resolve_frameset_sheets",
        "_normalize_meta_charset",
        "_clean_word_html",
    ):
        assert not hasattr(office_mod, name), f"{name} must be removed"


def test_convert_20mb_cap_on_intermediate(tmp_path):
    import mid.backends.office as office_mod

    src = tmp_path / "in.doc"
    src.write_text("x", encoding="utf-8")

    def big_factory(progid):
        app = MagicMock()

        def _saveas(out_path, *args, **kwargs):
            Path(out_path).write_bytes(b"x" * 100)

        docs = MagicMock()
        app.Documents.Open.return_value = docs
        docs.SaveAs.side_effect = _saveas
        return app

    with patch.object(sys, "platform", "win32"):
        with patch.object(office_mod, "_COM_FACTORY", big_factory):
            with patch("mid.backends.office._MAX_OUTPUT_BYTES", 10):
                r = office_mod.OfficeBackend().convert(src)
                assert r.success is False
                assert "20 MB" in (r.error or "") or "exceeds" in (r.error or "").lower()


def test_convert_delegates_ooxml_path_to_markitdown(tmp_path):
    """The OOXML intermediate path (not HTML) reaches MarkItDown."""
    import mid.backends.office as office_mod

    src = tmp_path / "in.doc"
    src.write_text("x", encoding="utf-8")
    seen = []
    with patch.object(sys, "platform", "win32"):
        with patch.object(office_mod, "_COM_FACTORY", lambda progid: _fake_com_success(progid, seen)):
            mock_md = MagicMock()
            mock_md.convert.return_value = MagicMock(success=True, content="# ok", error=None)
            with patch("mid.converters.markitdown.MarkitDownConverter", return_value=mock_md):
                r = office_mod.OfficeBackend().convert(src)
    assert r.success is True
    delegated = Path(mock_md.convert.call_args[0][0])
    assert delegated.suffix == ".docx"
    assert delegated.name == "in.docx"


def test_convert_multisheet_xlsx_all_sheets(tmp_path):
    """End-to-end with the REAL MarkItDown reader: every sheet converts.

    The fake COM layer writes a genuine 3-sheet .xlsx (as Excel SaveAs 51
    would); the backend must delegate it untouched so MarkItDown emits one
    ``## <name>`` section per sheet — no per-sheet join needed.
    """
    import mid.backends.office as office_mod

    src = tmp_path / "in.xls"
    src.write_text("x", encoding="utf-8")
    sheet_names = ["VersionHistory", "LibroBanco", "CUT"]

    def multisheet_factory(progid):
        from openpyxl import Workbook

        app = MagicMock()

        def _saveas(out_path, *args, **kwargs):
            book = Workbook()
            book.active.title = sheet_names[0]
            book.active["A1"] = "version"
            for name in sheet_names[1:]:
                ws = book.create_sheet(name)
                ws["A1"] = f"content-{name}"
            book.save(str(out_path))

        wb = MagicMock()
        app.Workbooks.Open.return_value = wb
        wb.SaveAs.side_effect = _saveas
        return app

    with patch.object(sys, "platform", "win32"):
        with patch.object(office_mod, "_COM_FACTORY", multisheet_factory):
            r = office_mod.OfficeBackend().convert(src)
    assert r.success is True
    for name in sheet_names:
        assert name in r.content, f"sheet {name} missing from output"
    assert r.content.count("## ") >= 3
    assert "content-LibroBanco" in r.content
    assert "content-CUT" in r.content


def test_convert_docx_end_to_end_real_markitdown(tmp_path):
    """End-to-end with the REAL MarkItDown reader: docx text flows through."""
    import mid.backends.office as office_mod

    src = tmp_path / "in.doc"
    src.write_text("x", encoding="utf-8")

    def rich_doc_factory(progid):
        from docx import Document

        app = MagicMock()

        def _saveas(out_path, *args, **kwargs):
            doc = Document()
            doc.add_heading("Objetivo", level=1)
            doc.add_paragraph("texto con ryas")
            doc.save(str(out_path))

        docs = MagicMock()
        app.Documents.Open.return_value = docs
        docs.SaveAs.side_effect = _saveas
        return app

    with patch.object(sys, "platform", "win32"):
        with patch.object(office_mod, "_COM_FACTORY", rich_doc_factory):
            r = office_mod.OfficeBackend().convert(src)
    assert r.success is True
    assert "Objetivo" in r.content
    assert "texto con ryas" in r.content
    assert "supportLists" not in r.content
    assert "\u00a0" not in r.content


def test_copy_exclusive_lock_maps_to_file_locked(tmp_path):
    """Exclusive lock on the pre-COM copy maps to file-locked, locale-independent (#BUG-2)."""
    import mid.backends.office as office_mod

    src = tmp_path / "in.xls"
    src.write_text("x", encoding="utf-8")

    def locked_copy(*args, **kwargs):
        exc = OSError("[WinError 32] El proceso no tiene acceso al archivo porque está siendo utilizado por otro proceso")
        exc.winerror = 32
        raise exc

    with patch.object(sys, "platform", "win32"):
        with patch("shutil.copy2", side_effect=locked_copy):
            with patch.object(office_mod, "_KILLER", MagicMock()) as mock_killer:
                r = office_mod.OfficeBackend().convert(src)
    assert r.success is False
    assert "file locked" in (r.error or "").lower()
    assert not mock_killer.called


def test_map_error_winerror_codes():
    """Numeric WinError codes map locale-independently (Spanish messages included)."""
    from mid.backends.office import _map_error

    for code in (32, 33):
        exc = OSError("cualquier mensaje localizado")
        exc.winerror = code
        assert "file locked" in _map_error(exc).lower()
    denied = OSError("Acceso denegado.")
    denied.winerror = 5
    assert "permission denied" in _map_error(denied).lower()


def test_convert_releases_app_reference_deterministically(tmp_path):
    """convert() must not retain the COM app wrapper after return (#29 no-orphan)."""
    import gc as _gc
    import weakref

    import mid.backends.office as office_mod

    src = tmp_path / "in.doc"
    src.write_text("x", encoding="utf-8")
    refs = []

    class Doc:
        def SaveAs(self, html_path, **kwargs):
            Path(html_path).write_text("<html><body>hi</body></html>", encoding="utf-8")

        def Close(self, *args):
            pass

    class Docs:
        def Open(self, *args, **kwargs):
            return Doc()

    class FakeApp:
        def __init__(self):
            self.Documents = Docs()

        def Quit(self):
            pass

    def factory(progid):
        app = FakeApp()
        refs.append(weakref.ref(app))
        return app

    with patch.object(sys, "platform", "win32"):
        with patch.object(office_mod, "_COM_FACTORY", factory):
            mock_md = MagicMock()
            mock_md.convert.return_value = MagicMock(success=True, content="# hi", error=None)
            with patch("mid.converters.markitdown.MarkitDownConverter", return_value=mock_md):
                r = office_mod.OfficeBackend().convert(src)
    assert r.success is True
    _gc.collect()
    assert refs[0]() is None, "COM app wrapper retained after convert() — orphan risk (#29)"


# --- OFFICE-05: error taxonomy --------------------------------------------------


def test_error_taxonomy_never_raises(tmp_path):
    """Locked/busy/orphan failures map to ConvertResult, never raise."""
    import mid.backends.office as office_mod

    src = tmp_path / "in.doc"
    src.write_text("x", encoding="utf-8")

    def locked_factory(progid):
        raise OSError("locked for editing by another user")

    with patch.object(sys, "platform", "win32"):
        with patch.object(office_mod, "_COM_FACTORY", locked_factory):
            r = office_mod.OfficeBackend().convert(src)
            assert r.success is False
            assert "close office" in (r.error or "").lower()

    def busy_factory(progid):
        raise RuntimeError("RPC_E_SERVERCALL_RETRYLATER call rejected")

    with patch.object(sys, "platform", "win32"):
        with patch.object(office_mod, "_COM_FACTORY", busy_factory):
            r2 = office_mod.OfficeBackend().convert(src)
            assert r2.success is False
            assert "busy" in (r2.error or "").lower() and "retry" in (r2.error or "").lower()


def test_convert_on_linux_reports_platform(tmp_path):
    """Linux convert short-circuits with unsupported-platform, never raises."""
    with patch.object(sys, "platform", "linux"):
        p = tmp_path / "a.doc"
        p.write_text("x", encoding="utf-8")
        r = _backend().convert(p)
        assert r.success is False
        assert "unsupported platform" in (r.error or "").lower()


# --- OOXML markdown normalization (NBSP + NaN + Unnamed) ------------------------


def test_normalize_nbsp_to_space_and_collapse():
    """Literal U+00A0 becomes a regular space; runs collapse, newlines kept."""
    from mid.backends.office import _normalize_ooxml_markdown

    out = _normalize_ooxml_markdown("a  b\nc  d")
    assert " " not in out
    assert out == "a b\nc d"


def test_normalize_whole_cell_nan_only():
    """Whole-cell NaN empties; words containing 'nan' (financiero) survive."""
    from mid.backends.office import _normalize_ooxml_markdown

    table = "| a | NaN | financiero |\n|---|---|---|\n| nan | NAN | medio de pago |"
    out = _normalize_ooxml_markdown(table)
    assert "financiero" in out
    assert "medio de pago" in out
    assert "NaN" not in out and "nan" not in out.replace("financiero", "")
    # column count stable: every row keeps 3 cells
    for line in out.splitlines():
        assert line.count("|") == 4
    # prose NaN without pipes is untouched
    assert _normalize_ooxml_markdown("NaN values indicate gaps") == "NaN values indicate gaps"


def test_normalize_unnamed_header_keeps_columns():
    """Exact 'Unnamed: N' headers blank; pipe count (columns) unchanged."""
    from mid.backends.office import _normalize_ooxml_markdown

    header = "| Unnamed: 0 | Name | Unnamed: 12 |"
    out = _normalize_ooxml_markdown(header)
    assert "Unnamed" not in out
    assert out.count("|") == header.count("|")
    # real header text is never blanked
    assert _normalize_ooxml_markdown("| Renamed: 3 | Name |") == "| Renamed: 3 | Name |"


def test_normalize_never_raises():
    """Normalizer is total: empty input passes through, never raises."""
    from mid.backends.office import _normalize_ooxml_markdown

    assert _normalize_ooxml_markdown("") == ""
    assert isinstance(_normalize_ooxml_markdown("| a | NaN |\n"), str)


# --- Merged-cell forward-fill (AI legibility) -------------------------------------


def _make_merged_fixture(path: Path) -> Path:
    """Matriz-like sheet: merged header zone + vertical merges + spacer row."""
    from openpyxl import Workbook

    book = Workbook()
    ws = book.active
    ws.title = "Banco"
    ws["A1"] = "Transición"
    ws.merge_cells("A1:A2")  # vertical header merge
    ws["B1"] = "Condiciones"
    ws.merge_cells("B1:C1")  # horizontal header merge
    ws["D1"] = "Impactos"
    ws.merge_cells("D1:E1")  # horizontal header merge
    ws["B2"] = "Cond 1"
    ws["C2"] = "Cond n"
    ws["D2"] = "Código"
    ws["E2"] = "Tipo"
    ws["A3"] = "Confirmar"
    ws.merge_cells("A3:A4")  # vertical body merge
    ws["B3"] = "x1"
    ws["C3"] = "y1"
    ws["D3"] = "m1"
    ws["E3"] = "t1"
    ws["B4"] = "x2"
    ws["C4"] = "y2"
    ws["D4"] = "m2"
    # E4 genuinely empty (never merged) — must NOT be filled.
    # Row 5 all-empty spacer row (never merged) — must stay empty.
    ws["A6"] = "Otro"
    ws["B6"] = "z"
    book.save(str(path))
    return path


def test_expand_merged_cells_horizontal_and_vertical(tmp_path):
    """2x2-style merges fill both axes; real empties stay empty."""
    from openpyxl import load_workbook

    from mid.backends.office import _expand_merged_cells

    p = _make_merged_fixture(tmp_path / "m.xlsx")
    assert _expand_merged_cells(p) is None
    ws = load_workbook(p)["Banco"]
    assert list(ws.merged_cells.ranges) == []
    # vertical fills
    assert ws["A2"].value == "Transición"
    assert ws["A4"].value == "Confirmar"
    # horizontal fills
    assert ws["C1"].value == "Condiciones"
    assert ws["E1"].value == "Impactos"
    # genuine empties untouched: spacer row + never-merged E4
    assert ws["A5"].value is None and ws["C5"].value is None
    assert ws["E4"].value is None


def test_expand_merged_cells_never_raises(tmp_path):
    """Missing file and missing openpyxl both fall back silently."""
    import builtins

    from mid.backends import office as office_mod

    assert office_mod._expand_merged_cells(tmp_path / "missing.xlsx") is None
    p = _make_merged_fixture(tmp_path / "m.xlsx")
    real_import = builtins.__import__

    def fake_import(name, *args, **kwargs):
        if name == "openpyxl":
            raise ImportError("No module named openpyxl")
        return real_import(name, *args, **kwargs)

    with patch.object(builtins, "__import__", fake_import):
        assert office_mod._expand_merged_cells(p) is None
    # fallback left the merges untouched
    from openpyxl import load_workbook

    assert len(list(load_workbook(p)["Banco"].merged_cells.ranges)) == 4


def test_convert_xls_merged_headers_forward_filled(tmp_path):
    """End-to-end with the REAL MarkItDown reader: merged fixture has no gaps.

    The fake COM layer copies a merged xlsx as the intermediate (as Excel
    SaveAs 51 preserving merges would); every output row must be
    self-contained while the spacer row stays empty.
    """
    import re

    import mid.backends.office as office_mod

    src = tmp_path / "in.xls"
    src.write_text("x", encoding="utf-8")
    fixture = _make_merged_fixture(tmp_path / "fixture.xlsx")

    def merged_factory(progid):
        import shutil

        app = MagicMock()

        def _saveas(out_path, *args, **kwargs):
            shutil.copy2(fixture, out_path)

        wb = MagicMock()
        app.Workbooks.Open.return_value = wb
        wb.SaveAs.side_effect = _saveas
        return app

    with patch.object(sys, "platform", "win32"):
        with patch.object(office_mod, "_COM_FACTORY", merged_factory):
            r = office_mod.OfficeBackend().convert(src)
    assert r.success is True
    assert "Unnamed" not in r.content
    assert not re.search(r"(?<=\|)\s*nan\s*(?=\|)", r.content, re.IGNORECASE)
    # forward-filled values repeat on every row they span
    assert r.content.count("Transición") >= 2
    assert r.content.count("Confirmar") >= 2
    # header row has no empty cells (pandas dedupes repeats as Name.1)
    header = next(line for line in r.content.splitlines() if "Condiciones" in line and "Cond 1" not in line)
    assert "||" not in header.replace(" ", "")
    # spacer row stays an all-empty separator, never filled with prior values
    assert any(re.fullmatch(r"[|\s]+", line) and line.count("|") >= 6 for line in r.content.splitlines())


# --- OFFICE-06: integration gate -------------------------------------------------


@pytest.mark.skipif(sys.platform != "win32", reason="requires Windows + installed Office")
def test_real_doc_conversion_win32(tmp_path):
    """Real .doc -> md on a licensed Windows runner (skipped on Linux)."""
    from mid.backends.office import OfficeBackend

    b = OfficeBackend()
    avail = b.probe()
    if not avail.available:
        pytest.skip(f"Office not detected on this runner: {avail.reason}")
    src = tmp_path / "sample.doc"
    src.write_bytes(b"placeholder")
    result = b.convert(src)
    assert result.success
    assert result.content.strip()
